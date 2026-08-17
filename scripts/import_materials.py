"""Convert the supplied TOEFL PDF/DOCX materials into deterministic JSON banks.

The source documents remain outside the web app.  Only structured study content is
written to data/, which keeps the first release private and avoids browser-side
document parsing.
"""

from __future__ import annotations

import json
import re
import unicodedata
from pathlib import Path

from docx import Document
from pypdf import PdfReader


ROOT = Path(__file__).resolve().parents[1]
DATA_DIR = ROOT / "data"
VOCABULARY_PDF = Path(r"D:\Desktop\英语\阅读\2026托福改革阅读高频必考800词（全话题完整版）-(1).pdf")
PHRASE_DOCX = Path(r"D:\Desktop\英语\阅读\【暑期配套资料】-托福阅读填词题与语法.docx")
ANSWER_DOCX = Path(r"D:\Desktop\英语\阅读\配套练习答案.docx")


def clean_text(value: str) -> str:
    value = unicodedata.normalize("NFKC", value)
    value = value.replace("\u00a0", " ")
    value = re.sub(r"中\s*文释义", "中文释义", value)
    value = re.sub(r"(?m)^\s*\d{1,3}/200\s*$", "", value)
    value = re.sub(r"[ \t]+", " ", value)
    value = re.sub(r" *\n *", "\n", value)
    return value.strip()


def one_line(value: str) -> str:
    return re.sub(r"\s+", " ", clean_text(value)).strip()


def find_example_form(word: str, example: str) -> str:
    tokens = re.findall(r"[A-Za-z]+(?:[-'][A-Za-z]+)*", example)
    lower_word = word.lower()

    if " " in lower_word:
        parts = lower_word.split()
        for start in range(0, len(tokens) - len(parts) + 1):
            window = tokens[start : start + len(parts)]
            if [token.lower() for token in window[:-1]] != parts[:-1]:
                continue
            last = window[-1].lower()
            expected = parts[-1]
            if last in {expected, expected + "s", expected + "es", expected + "ed", expected + "ing"}:
                return " ".join(window)

    exact = [token for token in tokens if token.lower() == lower_word]
    if exact:
        return exact[0]

    variants = {
        lower_word + "s",
        lower_word + "es",
        lower_word + "ed",
        lower_word + "ing",
        lower_word + "ly",
    }
    if lower_word.endswith("e"):
        variants.add(lower_word[:-1] + "ing")
        variants.add(lower_word + "d")
    if lower_word.endswith("y") and len(lower_word) > 2:
        variants.add(lower_word[:-1] + "ies")
        variants.add(lower_word[:-1] + "ied")
    for token in tokens:
        if token.lower() in variants:
            return token

    stem = lower_word[: max(4, min(7, len(lower_word) - 2))]
    for token in tokens:
        if token.lower().startswith(stem):
            return token

    if len(lower_word) >= 5:
        for token in tokens:
            if lower_word in token.lower():
                return token

    raise ValueError(f"Could not find {word!r} in example: {example}")


def make_example_cloze(word: str, example: str) -> dict[str, str | int]:
    answer = find_example_form(word, example)
    shown = max(1, int(len(answer) * 0.6))
    if len(answer) >= 6:
        shown = max(3, shown)
    shown = min(shown, len(answer) - 1)
    prefix = answer[:shown]
    prompt = re.sub(
        rf"\b{re.escape(answer)}\b",
        f"{prefix}  [{len(answer) - shown}]",
        example,
        count=1,
        flags=re.IGNORECASE,
    )
    return {"prompt": prompt, "answer": answer, "prefixLength": shown}


def parse_vocabulary() -> list[dict[str, object]]:
    reader = PdfReader(str(VOCABULARY_PDF))
    pages = [clean_text(page.extract_text() or "") for page in reader.pages]
    text = "\n".join(pages)
    header = re.compile(r"(?m)^(\d{1,3})\.\s+([A-Za-z][A-Za-z0-9 /\-'’]+?)\s*$")
    matches = list(header.finditer(text))
    candidates: dict[int, list[tuple[str, str]]] = {}
    for i, match in enumerate(matches):
        index = int(match.group(1))
        word = match.group(2).strip()
        end = matches[i + 1].start() if i + 1 < len(matches) else len(text)
        candidates.setdefault(index, []).append((word, clean_text(text[match.end() : end])))

    records: list[dict[str, object]] = []
    for index in range(1, 801):
        if index not in candidates:
            raise ValueError(f"Vocabulary item {index} is missing")
        word, segment = max(candidates[index], key=lambda item: len(item[1]))
        segment = re.sub(r"(?m)^\s*[一二三四五六七八九十]+、.*$", "", segment)

        pronunciation_match = re.search(r"(?m)^(/[^\n]+\(美式\))\s*$", segment)
        if not pronunciation_match:
            raise ValueError(f"Pronunciation missing for {index}. {word}")
        pronunciation = one_line(pronunciation_match.group(1))

        remainder = segment[pronunciation_match.end() :]
        definition_marker = remainder.find("英文释义:")
        if definition_marker < 0:
            raise ValueError(f"Definition missing for {index}. {word}")
        pos_meaning = one_line(remainder[:definition_marker])
        pos_match = re.match(r"^(n\.|v\.|adj\.|adv\.|prep\.|conj\.|pron\.|num\.|modal v\.|phr\.)\s*(.*)$", pos_meaning)
        if pos_match:
            labels = re.findall(r"(?<![A-Za-z])(n\.|v\.|adj\.|adv\.|prep\.|conj\.|pron\.|num\.|modal v\.|phr\.)", pos_meaning)
            part_of_speech = "/".join(dict.fromkeys(labels))
            meanings_zh = pos_match.group(2).strip()
        else:
            part_of_speech = ""
            meanings_zh = pos_meaning

        after_definition = remainder[definition_marker + len("英文释义:") :]
        example_marker = after_definition.find("托福阅读例句:")
        if example_marker < 0:
            raise ValueError(f"Example missing for {index}. {word}")
        before_example = after_definition[:example_marker]
        mnemonic_marker = before_example.find("记忆方法:")
        supplemental_marker = before_example.find("补充考点:")
        detail_markers = [marker for marker in (mnemonic_marker, supplemental_marker) if marker >= 0]
        detail_marker = min(detail_markers) if detail_markers else len(before_example)
        definition_block = one_line(before_example[:detail_marker])
        split_definition = re.split(r";\s*中文释义:\s*", definition_block, maxsplit=1)
        definition_en = split_definition[0].strip()
        notes_zh = split_definition[1].strip() if len(split_definition) > 1 else ""
        mnemonic = one_line(before_example[detail_marker:])
        mnemonic = re.sub(r"^(?:记忆方法|补充考点):\s*", "", mnemonic)
        example = one_line(after_definition[example_marker + len("托福阅读例句:") :])
        if index == 800:
            example = example.split("2026托福改革后阅读高频800词使用指南", 1)[0].strip()
        cloze = make_example_cloze(word, example)

        records.append(
            {
                "id": f"word-{index:03d}",
                "index": index,
                "word": word,
                "pronunciation": pronunciation,
                "partOfSpeech": part_of_speech,
                "meaningsZh": meanings_zh,
                "definitionEn": definition_en,
                "notesZh": notes_zh,
                "mnemonic": mnemonic,
                "example": example,
                "exampleCloze": cloze,
                "acceptedAnswers": [word],
            }
        )

    return records


def parse_phrases() -> list[dict[str, object]]:
    document = Document(str(PHRASE_DOCX))
    paragraphs = [one_line(p.text) for p in document.paragraphs]
    ranges = [range(44, 59), range(60, 70), range(71, 81), range(82, 97), range(100, 130), range(133, 153)]
    category_names = ["To 搭配", "With 搭配", "Of 搭配", "From / In / On 搭配", "结构与逻辑连接", "学术表达"]
    phrases: list[dict[str, object]] = []
    index = 0
    for category, paragraph_range in zip(category_names, ranges):
        for paragraph_index in paragraph_range:
            line = paragraphs[paragraph_index]
            if ":" not in line:
                raise ValueError(f"Phrase line has no separator: paragraph {paragraph_index}: {line}")
            english, meaning = line.split(":", 1)
            index += 1
            phrases.append(
                {
                    "id": f"phrase-{index:03d}",
                    "index": index,
                    "phrase": english.strip(),
                    "meaningZh": meaning.strip(),
                    "category": category,
                    "acceptedAnswers": [english.strip()],
                }
            )
    if len(phrases) != 100:
        raise ValueError(f"Expected 100 phrases, got {len(phrases)}")
    return phrases


def source_cloze_to_prompt(sentence: str) -> str:
    sentence = one_line(sentence)
    pattern = re.compile(r"([A-Za-z]+)?((?:_\s*)+)")

    def replace(match: re.Match[str]) -> str:
        prefix = match.group(1) or ""
        missing = match.group(2).count("_")
        following = match.string[match.end() : match.end() + 1]
        separator = " " if following and following[0].isalnum() else ""
        return f"{prefix}  [{missing}]{separator}"

    return pattern.sub(replace, sentence).strip()


def parse_exercises() -> list[dict[str, object]]:
    questions_doc = Document(str(PHRASE_DOCX))
    question_paragraphs = [one_line(p.text) for p in questions_doc.paragraphs]
    question_indexes = list(range(297, 347)) + list(range(349, 379)) + list(range(381, 401))
    questions = [question_paragraphs[index] for index in question_indexes]

    answer_doc = Document(str(ANSWER_DOCX))
    answer_paragraphs = [one_line(p.text) for p in answer_doc.paragraphs]
    answer_indexes = list(range(1, 51)) + list(range(53, 83)) + list(range(84, 104))
    answers = [answer_paragraphs[index] for index in answer_indexes]
    if len(questions) != 100 or len(answers) != 100:
        raise ValueError(f"Expected 100 questions/answers, got {len(questions)}/{len(answers)}")

    records: list[dict[str, object]] = []
    for i, (question, answer) in enumerate(zip(questions, answers), 1):
        explanation = ""
        canonical = answer
        if "(" in answer and ")" in answer:
            before, rest = answer.split("(", 1)
            inside, after = rest.split(")", 1)
            canonical = inside.strip() or before.strip()
            explanation = after.strip()
        elif i == 74 and "with respect to" in answer:
            canonical = "with respect to"
            explanation = answer[len("respect") :].strip()

        records.append(
            {
                "id": f"exercise-{i:03d}",
                "index": i,
                "sourceSentence": question,
                "prompt": source_cloze_to_prompt(question),
                "answer": canonical,
                "acceptedAnswers": list(dict.fromkeys(filter(None, [canonical, before.strip() if "(" in answer else ""]))),
                "sourceAnswer": answer,
                "explanation": explanation,
                "section": "介词搭配" if i <= 50 else "逻辑连接" if i <= 80 else "学术表达",
            }
        )
    return records


def write_json(name: str, payload: object) -> None:
    DATA_DIR.mkdir(parents=True, exist_ok=True)
    target = DATA_DIR / name
    target.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")


def main() -> None:
    vocabulary = parse_vocabulary()
    phrases = parse_phrases()
    exercises = parse_exercises()
    write_json("vocabulary.json", vocabulary)
    write_json("phrases.json", phrases)
    write_json("phrase-exercises.json", exercises)
    print(f"Imported {len(vocabulary)} vocabulary words, {len(phrases)} phrases, and {len(exercises)} exercises.")
    print("Vocabulary spot-check:", [(vocabulary[i - 1]["index"], vocabulary[i - 1]["word"]) for i in (1, 16, 483, 673, 800)])
    print("Phrase spot-check:", [(phrases[i - 1]["index"], phrases[i - 1]["phrase"]) for i in (1, 50, 51, 80, 100)])


if __name__ == "__main__":
    main()
